from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output-docx"
OUTPUT.mkdir(exist_ok=True)

NAVY = "132238"
ACCENT = "B8642F"
MUTED = "586476"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_paragraph_bidi(paragraph, enabled=True):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1" if enabled else "0")


def set_run_font(run, name="Arial", size=9, bold=False, color=NAVY):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc, rtl=False):
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.05

    for style_name in ["Title", "Subtitle", "Section Heading", "Role Heading", "Meta"]:
        if style_name not in doc.styles:
            doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(0)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(11.5)
    subtitle.font.bold = True
    subtitle.font.color.rgb = RGBColor.from_string(ACCENT)
    subtitle.paragraph_format.space_after = Pt(6)

    heading = doc.styles["Section Heading"]
    heading.font.name = "Arial"
    heading.font.size = Pt(11)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(NAVY)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), ACCENT)
    borders.append(bottom)
    heading._element.get_or_add_pPr().append(borders)

    role = doc.styles["Role Heading"]
    role.font.name = "Arial"
    role.font.size = Pt(9.8)
    role.font.bold = True
    role.font.color.rgb = RGBColor.from_string(NAVY)
    role.paragraph_format.space_before = Pt(4)
    role.paragraph_format.space_after = Pt(0)
    role.paragraph_format.keep_with_next = True

    meta = doc.styles["Meta"]
    meta.font.name = "Arial"
    meta.font.size = Pt(8.2)
    meta.font.color.rgb = RGBColor.from_string(MUTED)
    meta.paragraph_format.space_after = Pt(2)
    meta.paragraph_format.keep_with_next = True

    if rtl:
        section.top_margin = Inches(0.34)
        section.bottom_margin = Inches(0.32)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)
        styles = [normal, title, subtitle, heading, role, meta]
        for style in styles:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        normal.paragraph_format.space_after = Pt(1.2)
        normal.paragraph_format.line_spacing = 0.96
        heading.paragraph_format.space_before = Pt(5)
        heading.paragraph_format.space_after = Pt(2.5)
        role.paragraph_format.space_before = Pt(2)
        meta.paragraph_format.space_after = Pt(1)


def add_paragraph(doc, text, style=None, rtl=False, size=None, color=None, bold=False, space_after=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    if rtl:
        set_paragraph_bidi(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size or 9.3, color=color or NAVY, bold=bold)
    if space_after is not None:
        paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def add_bullet(doc, text, rtl=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0 if rtl else 0.14)
    paragraph.paragraph_format.right_indent = Inches(0.14 if rtl else 0)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12 if not rtl else 0)
    paragraph.paragraph_format.space_after = Pt(0.6 if rtl else 1.4)
    paragraph.paragraph_format.line_spacing = 0.96 if rtl else 1.02
    if rtl:
        set_paragraph_bidi(paragraph)
        prefix = "• "
    else:
        prefix = "• "
    run = paragraph.add_run(prefix + text)
    set_run_font(run, size=7.9 if rtl else 8.8)


def add_header(doc, data, rtl=False):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    if rtl:
        table.columns[0].width = Inches(2.15)
        table.columns[1].width = Inches(4.95)
        main_cell, contact_cell = table.cell(0, 1), table.cell(0, 0)
    else:
        table.columns[0].width = Inches(4.95)
        table.columns[1].width = Inches(2.15)
        main_cell, contact_cell = table.cell(0, 0), table.cell(0, 1)

    for cell in [main_cell, contact_cell]:
        cell.vertical_alignment = 1
        cell.margin_top = 0
        cell.margin_bottom = 0
        cell.margin_left = 0
        cell.margin_right = 0

    p = main_cell.paragraphs[0]
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    if rtl:
        set_paragraph_bidi(p)
    set_run_font(p.add_run(data["name"]), size=24, bold=True)

    p = main_cell.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    if rtl:
        set_paragraph_bidi(p)
    set_run_font(p.add_run(data["title"]), size=11.5, bold=True, color=ACCENT)

    p = contact_cell.paragraphs[0]
    first_is_arabic = any("\u0600" <= char <= "\u06ff" for char in data["contact"][0])
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if first_is_arabic else WD_ALIGN_PARAGRAPH.LEFT
    if rtl and first_is_arabic:
        set_paragraph_bidi(p)
    for index, line in enumerate(data["contact"]):
        if index:
            p = contact_cell.add_paragraph()
            is_arabic = any("\u0600" <= char <= "\u06ff" for char in line)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if is_arabic else WD_ALIGN_PARAGRAPH.LEFT
            if rtl and is_arabic:
                set_paragraph_bidi(p)
        p.paragraph_format.space_after = Pt(1)
        set_run_font(p.add_run(line), size=8, color=MUTED)


def add_section(doc, title, rtl=False):
    return add_paragraph(doc, title, style="Section Heading", rtl=rtl, size=10 if rtl else 11, bold=True)


def add_experience(doc, item, rtl=False):
    add_paragraph(doc, f'{item["role"]} | {item["company"]}', style="Role Heading", rtl=rtl, size=9 if rtl else 9.8, bold=True)
    add_paragraph(doc, f'{item["place"]} | {item["date"]}', style="Meta", rtl=rtl, size=7.5 if rtl else 8.2, color=MUTED)
    for bullet in item["bullets"]:
        add_bullet(doc, bullet, rtl=rtl)


def add_project(doc, item, rtl=False):
    add_paragraph(doc, item["title"], style="Role Heading", rtl=rtl, size=9 if rtl else 9.8, bold=True)
    add_paragraph(doc, item["stack"], style="Meta", rtl=rtl, size=7.5 if rtl else 8.2, color=ACCENT)
    add_bullet(doc, item["description"], rtl=rtl)
    add_paragraph(doc, item["link"], style="Meta", rtl=rtl, size=7.3 if rtl else 7.8, color=MUTED)


def build_resume(data, filename, rtl=False):
    doc = Document()
    configure_document(doc, rtl=rtl)
    add_header(doc, data, rtl=rtl)
    add_paragraph(doc, data["summary"], rtl=rtl, size=8.4 if rtl else 9.4, space_after=2 if rtl else 4)

    add_section(doc, data["experience_heading"], rtl=rtl)
    for item in data["experience"]:
        add_experience(doc, item, rtl=rtl)

    add_section(doc, data["projects_heading"], rtl=rtl)
    for project in data["projects"]:
        add_project(doc, project, rtl=rtl)

    add_section(doc, data["education_heading"], rtl=rtl)
    add_paragraph(doc, data["degree"], style="Role Heading", rtl=rtl, size=9 if rtl else 9.8, bold=True)
    add_paragraph(doc, data["education_meta"], style="Meta", rtl=rtl, size=7.5 if rtl else 8.2, color=MUTED)

    add_section(doc, data["academic_heading"], rtl=rtl)
    for item in data["academic"]:
        add_bullet(doc, item, rtl=rtl)

    add_section(doc, data["skills_heading"], rtl=rtl)
    for label, values in data["skills"]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
        if rtl:
            set_paragraph_bidi(paragraph)
        set_run_font(paragraph.add_run(label + ": "), size=7.9 if rtl else 8.8, bold=True, color=ACCENT)
        set_run_font(paragraph.add_run(values), size=7.9 if rtl else 8.8)
        paragraph.paragraph_format.space_after = Pt(2)

    path = OUTPUT / filename
    doc.save(path)
    return path


EN = {
    "name": "Abdelaziz Abuthuraya",
    "title": "Industrial Engineer & Digital Systems Developer",
    "contact": ["Jeddah, Saudi Arabia", "azizkza99@gmail.com", "+966 54 994 9023", "linkedin.com/in/abdelazizabuthuraya", "github.com/azizkza99"],
    "summary": "Industrial Engineering graduate with hands-on experience in production planning, process improvement, quality, and supply-chain analysis across manufacturing and EPC environments in Saudi Arabia and Türkiye. Combines engineering thinking with Arabic-first web systems built using React, TypeScript, Supabase, and Three.js. Trilingual in Arabic, English, and Turkish.",
    "experience_heading": "PROFESSIONAL EXPERIENCE",
    "experience": [
        {"role": "Production Planning Intern", "company": "EDAS Yatak", "place": "Kayseri, Türkiye", "date": "Aug 2025 - Oct 2025", "bullets": ["Prepared and monitored daily production schedules, coordinating across departments to keep workflow aligned with delivery priorities.", "Reviewed production-performance data to identify bottlenecks and translate findings into practical workflow improvements.", "Applied production-planning and industrial-engineering principles to improve task sequencing and process organization."]},
        {"role": "Operations & Process Improvement Intern", "company": "Murat Şahin Orman Ürünleri", "place": "Tokat, Türkiye", "date": "Jan 2025 - Jul 2025", "bullets": ["Worked across plywood and MDF cutting, pressing, and finishing stages while observing handoffs and constraints along the production flow.", "Tracked daily output, packaging, and storage activities, highlighting opportunities to reduce waste and improve material movement.", "Performed surface and dimensional quality checks and used SolidWorks for basic engineering-support and modeling tasks."]},
        {"role": "Project Management Intern", "company": "Delta CJSC", "place": "Jeddah, Saudi Arabia", "date": "Jul 2024 - Aug 2024", "bullets": ["Organized project files and operational data while supporting daily EPC coordination and a clear view of work status.", "Prepared Excel reports for schedules and logistics and contributed ideas during workflow-improvement discussions."]},
        {"role": "Engineering Intern", "company": "Delta CJSC", "place": "Riyadh, Saudi Arabia", "date": "Jul 2022 - Aug 2022", "bullets": ["Supported coordination across engineering, procurement, and construction while monitoring timelines and project KPIs.", "Improved scheduling and logistics reports in Excel and contributed to discussions on more efficient service delivery."]},
    ],
    "projects_heading": "SELECTED TECHNICAL PROJECTS",
    "projects": [
        {"title": "FORM - 3D Product Configurator", "stack": "React | Three.js | Vite | WebGL", "description": "Translated industrial door specifications into a bilingual interactive experience with a procedural 3D model, live specification sheet, accessible controls, and WebGL fallback.", "link": "bezi-product-viewer.vercel.app | github.com/azizkza99/bezi-product-viewer"},
        {"title": "Kayan - Workflow Automation Concept", "stack": "React | TypeScript | Supabase | Row-Level Security", "description": "Built an Arabic-first workflow discovery experience with real Supabase submission, database access controls, localized validation, and responsive feedback states.", "link": "kayan-app-henna.vercel.app | github.com/azizkza99/kayan-app"},
        {"title": "ECLIPSE - Premium Coffee Experience", "stack": "TanStack Start | React | TypeScript | Resend", "description": "Developed a bilingual product experience with server-validated booking and a coffee recommendation flow; reduced video assets from about 79 MB to under 8 MB.", "link": "eclipse-luxury-experience.vercel.app"},
    ],
    "education_heading": "EDUCATION",
    "degree": "BSc in Industrial Engineering - Üsküdar University",
    "education_meta": "Istanbul, Türkiye | 2020 - 2024 | GPA: 2.75 / 4.00",
    "academic_heading": "ACADEMIC ENGINEERING PROJECTS",
    "academic": ["Workforce planning: designed and analysed a 14-question survey across eight companies covering selection, development, and succession planning.", "Supply chain and revenue: analysed 85% machine utilisation and waste rates of 0.5-1.5% to identify operational improvement opportunities.", "Virtual machine placement: built a QoS- and SLA-constrained MIP model that achieved a zero optimality gap in the tested scenarios."],
    "skills_heading": "SKILLS",
    "skills": [("Engineering", "Production planning, process improvement, Lean, 5S, quality control, supply chain, mathematical modelling"), ("Technology", "React, TypeScript, JavaScript, TanStack Start, Supabase/PostgreSQL, Three.js, HTML/CSS, Git/GitHub, Vercel"), ("Tools", "Microsoft Excel, SolidWorks, AutoCAD, Microsoft Office"), ("Languages", "Arabic - Native | English - C1 | Turkish - B1")],
}


AR = {
    "name": "عبد العزيز أبو ثريا",
    "title": "مهندس صناعي ومطور أنظمة رقمية",
    "contact": ["جدة، المملكة العربية السعودية", "azizkza99@gmail.com", "+966 54 994 9023", "linkedin.com/in/abdelazizabuthuraya", "github.com/azizkza99"],
    "summary": "مهندس صناعي بخبرة ميدانية في تخطيط الإنتاج، وتحسين العمليات، والجودة، وتحليل سلاسل الإمداد ضمن بيئات تصنيع ومشاريع في السعودية وتركيا. أدمج التفكير الهندسي مع تطوير أنظمة عربية أولًا باستخدام React وTypeScript وSupabase وThree.js. اللغات: العربية (الأم)، الإنجليزية (C1)، والتركية (B1).",
    "experience_heading": "الخبرة المهنية",
    "experience": [
        {"role": "متدرب تخطيط إنتاج", "company": "EDAS Yatak", "place": "قيصري، تركيا", "date": "أغسطس 2025 - أكتوبر 2025", "bullets": ["إعداد ومتابعة جداول الإنتاج اليومية والتنسيق بين الأقسام للحفاظ على تدفق العمل وربط الأولويات بمواعيد التسليم.", "تحليل بيانات الأداء لتحديد الاختناقات وتحويل النتائج إلى مقترحات عملية لتحسين تنظيم العمليات.", "تطبيق مبادئ تخطيط الإنتاج والهندسة الصناعية لتحسين تسلسل المهام ووضوح سير العمل."]},
        {"role": "متدرب عمليات وتحسين عمليات", "company": "Murat Şahin Orman Ürünleri", "place": "توكات، تركيا", "date": "يناير 2025 - يوليو 2025", "bullets": ["المشاركة ميدانيًا في مراحل قطع وكبس وتشطيب ألواح Plywood وMDF ومراقبة نقاط التسليم بين مراحل الإنتاج.", "متابعة أهداف الإنتاج والتعبئة والتخزين، مع رصد فرص تقليل الهدر وتحسين حركة المواد والمنتجات الجاهزة.", "تنفيذ فحوصات جودة سطحية وبُعدية واستخدام SolidWorks في مهام دعم هندسي ونمذجة أساسية."]},
        {"role": "متدرب إدارة مشاريع", "company": "Delta CJSC", "place": "جدة، السعودية", "date": "يوليو 2024 - أغسطس 2024", "bullets": ["تنظيم ملفات المشروع وبياناته ودعم التنسيق اليومي لفريق EPC للحفاظ على وضوح حالة الأعمال.", "إعداد تقارير Excel للجداول الزمنية واللوجستيات والمشاركة في اجتماعات تحسين كفاءة العمل."]},
        {"role": "متدرب هندسي", "company": "Delta CJSC", "place": "الرياض، السعودية", "date": "يوليو 2022 - أغسطس 2022", "bullets": ["دعم التنسيق بين الهندسة والمشتريات والإنشاءات ومتابعة الجداول الزمنية ومؤشرات الأداء للمشروع.", "تحسين تقارير الجدولة واللوجستيات باستخدام Excel والمساهمة في مناقشات تطوير سير العمل."]},
    ],
    "projects_heading": "مشاريع تقنية مختارة",
    "projects": [
        {"title": "FORM - أداة تهيئة منتج ثلاثية الأبعاد", "stack": "React | Three.js | Vite | WebGL", "description": "تحويل مواصفات باب صناعي إلى تجربة تفاعلية ثنائية اللغة تشمل نموذجًا ثلاثي الأبعاد وورقة مواصفات فورية وتحكمًا متاحًا ودعمًا للأجهزة التي لا تشغّل WebGL.", "link": "bezi-product-viewer.vercel.app | github.com/azizkza99/bezi-product-viewer"},
        {"title": "Kayan - تصور لأتمتة سير العمل", "stack": "React | TypeScript | Supabase | Row-Level Security", "description": "بناء تجربة عربية أولًا لاكتشاف العمليات مع إرسال فعلي عبر Supabase وضوابط وصول للبيانات وتحقق محلي وحالات استجابة واضحة.", "link": "kayan-app-henna.vercel.app | github.com/azizkza99/kayan-app"},
        {"title": "ECLIPSE - تجربة قهوة مختصة فاخرة", "stack": "TanStack Start | React | TypeScript | Resend", "description": "تطوير تجربة ثنائية اللغة مع حجز متحقق منه على الخادم وأداة توصية للقهوة، وخفض وسائط الفيديو من نحو 79MB إلى أقل من 8MB.", "link": "eclipse-luxury-experience.vercel.app"},
    ],
    "education_heading": "التعليم",
    "degree": "بكالوريوس الهندسة الصناعية - Üsküdar University",
    "education_meta": "إسطنبول، تركيا | 2020 - 2024 | المعدل: 2.75 / 4.00",
    "academic_heading": "مشاريع هندسية أكاديمية",
    "academic": ["تخطيط الموارد البشرية: تصميم وتحليل استبيان من 14 سؤالًا عبر 8 شركات لدراسة الاختيار والتطوير والتعاقب الوظيفي.", "سلسلة الإمداد والإيرادات: تحليل استخدام الآلات بنسبة 85% ومعدلات هدر بين 0.5% و1.5% لتحديد فرص التحسين.", "توزيع الأجهزة الافتراضية: بناء نموذج MIP تحت قيود QoS وSLA وتحقيق حلول مثالية بفجوة مثالية صفرية في السيناريوهات المختبرة."],
    "skills_heading": "المهارات",
    "skills": [("الهندسة", "تخطيط الإنتاج، تحسين العمليات، Lean، 5S، الجودة، سلاسل الإمداد، النمذجة الرياضية، Excel، SolidWorks"), ("التقنية", "React، TypeScript، JavaScript، TanStack Start، Supabase/PostgreSQL، Three.js، HTML/CSS، Git/GitHub، Vercel")],
}


if __name__ == "__main__":
    print(build_resume(EN, "Abdelaziz-Abuthuraya-CV-EN.docx"))
    print(build_resume(AR, "Abdelaziz-Abuthuraya-CV-AR.docx", rtl=True))
